from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.dml import MSO_THEME_COLOR
from PIL import Image, ImageDraw, ImageFont
import qrcode, os, textwrap, math
from pathlib import Path

# The existing PowerPoint is the artwork source, not an output template.
# Extract its original picture backgrounds before overwriting the deck.
from io import BytesIO
BASE=Path(__file__).resolve().parent.parent
ROOT=BASE/'presentation'/'_oct10_build'
BG=ROOT/'bgs'
OUT=ROOT
BG.mkdir(parents=True,exist_ok=True)
old_deck=BASE/'presentation'/'Beyond-Chatbots-October10.pptx'
if not old_deck.exists():
    raise RuntimeError('Original presentation is missing: '+str(old_deck))
original=Presentation(str(old_deck))
if len(original.slides)<20:
    raise RuntimeError('Expected >=20 slides in the original artwork source')
def source_bg(i):
    slide=original.slides[i]
    pictures=[s for s in slide.shapes if s.shape_type==13]
    if not pictures:
        raise RuntimeError(f'No full-bleed artwork picture on original slide {i+1}')
    return pictures[0].image.blob
for i in range(20):
    png=source_bg(i)
    pic=Image.open(BytesIO(png)).convert('RGB')
    pic.resize((1280,720),Image.Resampling.LANCZOS).save(BG/f'bg{i+1:02d}.jpg','JPEG',quality=79,optimize=True)
(ROOT/'slide1_bg.png').write_bytes(source_bg(0))
qr_slide=next((i for i,slide in enumerate(original.slides) if any(
    'Questions & stay connected' in (getattr(sh,'text','') or '') for sh in slide.shapes
)),None)
if qr_slide is None:
    raise RuntimeError('Cannot locate the original QR/contact slide to retain its background')
qr_picture=source_bg(qr_slide)
(ROOT/'slide16_bg.png').write_bytes(qr_picture)
Image.open(BytesIO(qr_picture)).convert('RGB').resize((1280,720),Image.Resampling.LANCZOS).save(
    BG/'bg16.jpg','JPEG',quality=79,optimize=True
)
# New deck now builds locally from preserved original artwork.

W=13.333; H=7.5
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
BLANK=prs.slide_layouts[6]
C={
 'navy':'061524','ink':'051320','glass':'031321','cyan':'3DE8FF','aqua':'23D3D3','purple':'A17AFF',
 'white':'F4FBFF','muted':'AFC7DB','gold':'FFCE70','green':'6DF2A2','red':'FF9A9A','rule':'2D5774',
 'mid':'214D75','pale':'C5E9FF',
}

def rgb(c):
 c=C.get(c,c).strip('#')
 return RGBColor.from_string(c)

def alpha(shape, opacity=85):
 # DrawingML alpha is opacity, not transparency. 100% = 100000.
 el=shape.fill._xPr
 sf=el.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
 if sf is None: return
 color = sf.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
 if color is not None:
  child=OxmlElement('a:alpha');child.set('val',str(round(opacity*1000)))
  color.append(child)

def shape(sl,x,y,w,h,color='glass',opacity=85,line_color=None,line_width=1.0,radius=True):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color);alpha(sh,opacity)
 if line_color:
  sh.line.color.rgb=rgb(line_color);sh.line.width=Pt(line_width)
 else:sh.line.fill.background()
 return sh

def add_text(sl,x,y,w,h,text,size=23,color='white',bold=False,align='left',font='Aptos',margin=0.03,vert=MSO_ANCHOR.MIDDLE,paragraph_spacing=1.0):
 tx=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=tx.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(margin)
 tf.margin_top=tf.margin_bottom=Inches(margin)
 tf.vertical_anchor=vert
 p=tf.paragraphs[0]
 p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
 if '\n' in text:
  for j,line in enumerate(text.split('\n')):
   if j==0:
    run=p.add_run();run.text=line
   else:
    p=tf.add_paragraph();p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
    run=p.add_run();run.text=line
   run.font.name=font;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=rgb(color)
 else:
  run=p.add_run();run.text=text;run.font.name=font;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=rgb(color)
 for p in tf.paragraphs:p.space_after=Pt(paragraph_spacing)
 return tx

def picture_bg(sl,index,brightness=35):
 bg=BG/f'bg{index:02}.jpg'
 sl.shapes.add_picture(str(bg),0,0,width=Inches(W),height=Inches(H))
 if brightness:shape(sl,0,0,W,H,color='ink',opacity=brightness,radius=False)

def header(sl,index,title,section,source=''):
 shape(sl,.48,.29,12.35,.95,'navy',83,'rule',.8)
 add_text(sl,.68,.34,7.5,.29,'KIRA LABS  /  '+section,13,'cyan',True)
 add_text(sl,.65,.68,11.8,.52,title,31,'white',True)
 add_text(sl,12.04,.41,.57,.31,str(index).zfill(2),14,'aqua',True,'right')
 # footer
 shape(sl,0,7.00,W,.5,'ink',87,radius=False)
 add_text(sl,.66,7.105,8,.15,'BEYOND CHATBOTS  •  OCTOBER 10, 2026',8.0,'muted')
 add_text(sl,11.15,7.105,1.55,.15,'KIRALABS.ORG',8.0,'aqua',True,'right')
 if source:add_text(sl,.65,6.81,12.0,.15,source,7.5,'muted')

def slide(index,title,section,bg,brightness=44,source=''):
 s=prs.slides.add_slide(BLANK);picture_bg(s,bg,brightness);header(s,index,title,section,source);return s

def panel(sl,x,y,w,h,title,body,accent='cyan',body_size=20,title_size=21):
 shape(sl,x,y,w,h,'navy',91,accent,.95)
 add_text(sl,x+.20,y+.18,w-.4,.36,title,title_size,accent,True)
 add_text(sl,x+.20,y+.60,w-.4,h-.72,body,body_size,'white',False,vert=MSO_ANCHOR.TOP)

def pill(sl,x,y,w,label,col='cyan',font=17,h=.41):
 shape(sl,x,y,w,h,'ink',91,col,.7)
 add_text(sl,x+.03,y+.02,w-.06,h-.04,label,font,'white',True,'center')

def circle(sl,x,y,d,color,opacity=100):
 sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);alpha(sh,opacity);sh.line.fill.background()
 return sh

def rule(sl,x,y,w,color='cyan',height=.026):
 shape(sl,x,y,w,height,color,91,radius=False)

def qr(url,filename):
 obj=qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_H,box_size=14,border=4)
 obj.add_data(url);obj.make(fit=True);im=obj.make_image(fill_color='black',back_color='white').convert('RGB');im.save(ROOT/filename)
 return str(ROOT/filename)
repo='https://github.com/rmcmurrer81/Beyond-Chatbots'
site='https://kiralabs.org'
linked='https://www.linkedin.com/in/rmcmurrer'
qr_repo=qr(repo,'qr_beyond_chatbots.png');qr_site=qr(site,'qr_kiralabs.png');qr_linked=qr(linked,'qr_linkedin.png')

slides=[]
def notes(sl,t0,t1,script,extra=''):
 s=f'Slide {len(prs.slides)}  |  Elapsed {t0}–{t1}\n\n'+script
 if extra:s+='\n\nSPEAKER NOTE: '+extra
 sl.notes_slide.notes_text_frame.text=s
 slides.append({'number':len(prs.slides),'title':sl.shapes[2].text if False else '', 'start':t0,'end':t1,'script':script,'extra':extra})

# SLIDE 1 — exact existing cover background and footer matching source
s=prs.slides.add_slide(BLANK)
s.shapes.add_picture(str(ROOT/'slide1_bg.png'),0,0,width=Inches(W),height=Inches(H))
shape(s,0,6.55,W,.95,'ink',61,radius=False)
add_text(s,.72,6.88,5.1,.22,'Robert McMurrer • Kira Labs • October 10, 2026',9.5,'white')
add_text(s,9.15,6.88,3.2,.22,'Decentralized AI Day New York',9.5,'cyan',align='right')
add_text(s,.6,7.08,8.2,.18,'Early-stage research • local-first identity • measurable boundaries',6.6,'muted')
add_text(s,11.6,7.08,1.1,.18,'KiraLabs.org',6.6,'cyan',align='right')
notes(s,'0:00','0:35',"Good afternoon. I'm the founder of Kira Labs in Newark, New Jersey. Today I want to do something a little different. Before I describe NewBrain, I want to show how we got from the earliest ideas about thinking machines to the questions I'm testing today. Then I'll show the actual code and results I've made available. I'm not here to tell you Maya is finished. I want to show where the work really stands.")

# SLIDE 2 — EARLY QR, repo name enormous
s=slide(2,'BEYOND-CHATBOTS','OPEN THE REPOSITORY',16,brightness=47)
add_text(s,.82,1.31,11.7,.46,'SCAN NOW • KEEP THE RESEARCH AFTER THE TALK',22,'cyan',True)
shape(s,.77,1.88,5.75,4.76,'navy',95,'cyan',1.4)
s.shapes.add_picture(qr_repo,Inches(1.09),Inches(2.32),width=Inches(2.96),height=Inches(2.96))
add_text(s,4.18,2.63,2.11,.48,'GITHUB',24,'cyan',True)
add_text(s,4.18,3.10,2.15,2.00,'Code\nAI history PDF\nSpeaker packet\nResearch reports',18.5,'white')
add_text(s,1.09,5.47,5.05,.63,'github.com/rmcmurrer81/Beyond-Chatbots',15.5,'white',True)
shape(s,6.78,1.88,2.73,4.76,'navy',95,'purple',1.2)
s.shapes.add_picture(qr_site,Inches(7.20),Inches(2.60),width=Inches(1.91),height=Inches(1.91))
add_text(s,7.04,4.85,2.23,.45,'Kira Labs',21,'white',True,'center')
add_text(s,7.05,5.39,2.22,.42,'kiralabs.org',14.5,'muted',False,'center')
shape(s,9.75,1.88,2.73,4.76,'navy',95,'aqua',1.2)
s.shapes.add_picture(qr_linked,Inches(10.18),Inches(2.60),width=Inches(1.91),height=Inches(1.91))
add_text(s,10.03,4.85,2.16,.45,'LinkedIn',21,'white',True,'center')
add_text(s,9.99,5.39,2.22,.42,'Connect after',14.5,'muted',False,'center')
notes(s,'0:35','1:15',"Before we begin, this is the public GitHub repository: Beyond-Chatbots. Please scan the large code now so you have the link. It includes a 42-page illustrated AI history, my slides and speaking materials, NewBrain source references, Avatar Builder research, and copies of Aster, IdeaForge, and the earlier character chat software. You can browse while I speak, or use Code, Download ZIP at home. These are research copies: read their status notes before trying to run anything. The other codes go to Kira Labs and LinkedIn.","Pause for a few seconds while people scan the large QR. It points to the repository, not the GitHub profile.")

# SLIDE 3 History landmarks
s=slide(3,'The idea of artificial intelligence is old','HISTORY / FOUNDATIONS',2,45)
add_text(s,.83,1.42,11.35,.48,'Ideas came before computers. Computers made them testable.',23,'white',True)
found=[('1830s','Babbage','Programmable\nmachine'),('1843','Lovelace','Symbols beyond\narithmetic'),('1936','Turing','A general model\nof computation'),('1940s','Computing','Electronic\nexperiments')]
for i,(yr,name,txt) in enumerate(found):
 x=.75+i*3.15
 shape(s,x,2.19,2.92,2.64,'navy',92, ['cyan','purple','aqua','gold'][i],1.1)
 add_text(s,x+.17,2.35,2.57,.50,yr,27,['cyan','purple','aqua','gold'][i],True,'center')
 add_text(s,x+.12,3.05,2.66,.45,name,23,'white',True,'center')
 add_text(s,x+.18,3.74,2.56,.80,txt,18,'muted',False,'center')
shape(s,.88,5.33,11.58,.77,'navy',88,'rule',1)
add_text(s,1.15,5.47,11.1,.49,'The recurring question: can a machine do more than follow instructions?',21,'white',True,'center')
notes(s,'1:15','2:25',"Artificial intelligence didn't begin in 2022, or even in the 1950s. People imagined artificial beings for centuries. In the 1800s Charles Babbage designed the Analytical Engine, while Ada Lovelace argued that a machine might manipulate symbols, not only calculate numbers. In 1936 Alan Turing described a mathematical model of general computation. Electronic computers then made it possible to actually run increasingly complex programs. Notice that none of this yet proves machine consciousness. It establishes the machinery and ideas that made later experiments possible. The bigger question was always: can these machines reason, learn, or eventually become more than programmed tools?")

# SLIDE 4 Turing
s=slide(4,"Alan Turing's challenge",'HISTORY / 1950',3,51,'Source: A. Turing, Computing Machinery and Intelligence (1950).')
shape(s,.85,1.65,6.8,3.85,'navy',92,'cyan',1.1)
add_text(s,1.22,2.02,5.99,.63,'“Can machines think?”',37,'white',True)
add_text(s,1.23,2.95,5.85,1.67,'Turing proposed the imitation game: judge a machine through its conversational performance.',22,'pale')
shape(s,8.05,1.67,4.38,3.84,'navy',93,'purple',1.1)
add_text(s,8.48,1.97,3.5,.44,'WHAT IT TESTS',19,'cyan',True)
add_text(s,8.44,2.60,3.64,.73,'Observable dialogue',23,'white',True)
rule(s,8.45,3.51,3.37,'rule')
add_text(s,8.44,3.69,3.65,1.26,'It does NOT tell us everything about memory, emotion, or inner experience.',18.5,'muted')
notes(s,'2:25','3:40',"In 1950 Alan Turing published Computing Machinery and Intelligence. Instead of becoming stuck on the definition of thinking, he described the imitation game, which asked us to evaluate a machine by how convincingly it could participate in a conversation. That became the popular Turing Test. It's a historically important question, and conversation is a valuable skill. But it has limits. A system can say something convincing and still get facts wrong, forget what happened yesterday, or confuse one person's history with another. Turing's idea helped focus attention on observable behavior; it didn't solve the separate questions of experience, conscience, or genuine long-term learning.")

# SLIDE 5 early 1950s-60s
s=slide(5,'The first celebrated AI programs','HISTORY / REASONING',4,47)
three=[('1956','DARTMOUTH','Artificial intelligence becomes a named research field.'),('1956','LOGIC THEORIST','A computer searches for mathematical proofs.'),('1957+','GENERAL PROBLEM SOLVER','Researchers pursue reusable symbolic problem solving.')]
for i,(year,title,body) in enumerate(three):
 x=.68+i*4.25;shape(s,x,1.82,3.92,3.84,'navy',92,['cyan','purple','aqua'][i],1.15)
 add_text(s,x+.25,2.13,3.41,.5,year,26,['cyan','purple','aqua'][i],True,'center')
 add_text(s,x+.21,2.94,3.47,.69,title,23,'white',True,'center')
 add_text(s,x+.32,3.98,3.28,1.30,body,19,'muted',False,'center')
add_text(s,1.10,5.99,11.4,.48,'These were narrow demonstrations, not complete human minds.',21,'white',True,'center')
notes(s,'3:40','4:40',"The name artificial intelligence became established around the Dartmouth project in 1956. Around that time, Newell, Simon, and Shaw developed Logic Theorist, which could find proofs in a defined mathematical domain. General Problem Solver followed, exploring whether similar symbolic methods could work across different tasks. These were major accomplishments. They showed that computers could perform operations associated with reasoning. But they also showed the limitation of narrow tasks: solving a theorem doesn't automatically give a computer common sense, relationships, or an understanding of the world. You can see a much longer version of this history in the PDF in Beyond-Chatbots.")

# SLIDE 6 ELIZA
s=slide(6,'ELIZA: conversation before modern AI','HISTORY / 1966',5,45,'Source: Joseph Weizenbaum, ELIZA (1966). Dialogue below is illustrative.')
shape(s,.77,1.74,7.13,4.48,'ink',97,'aqua',1.2)
add_text(s,1.15,1.98,6.34,.4,'ELIZA / DOCTOR SCRIPT',22,'green',True,font='Consolas')
add_text(s,1.15,2.74,6.28,2.74,'PERSON:  I feel worried.\n\nELIZA:   Why do you feel worried?\n\nPERSON:  I think nobody listens.\n\nELIZA:   Tell me more about that.',18.8,'green',False,font='Consolas',vert=MSO_ANCHOR.TOP)
shape(s,8.22,1.75,4.34,4.46,'navy',93,'purple',1.1)
add_text(s,8.56,2.13,3.74,.39,'HOW IT WORKED',21,'cyan',True)
add_text(s,8.55,2.83,3.70,2.40,'Pattern matching\nScripted transformations\nReflected user wording\nNo general language model',19.8,'white')
notes(s,'4:40','5:50',"Now we get to ELIZA, created by Joseph Weizenbaum at MIT in the 1960s. ELIZA's famous DOCTOR script resembled a therapist. It could ask follow-up questions or reflect the user's own words back at them. The exchange shown here is an illustration, not an authentic 1966 transcript. The program did not have today's large neural models. It largely relied on rules and pattern matching. But to the person using it, those answers could feel surprisingly personal. Some people felt they were being listened to, even though the program wasn't performing the kind of deep comprehension they attributed to it. That's what makes ELIZA so important to this discussion.")

# SLIDE 7 ELIZA Effect
s=slide(7,'The ELIZA effect','HISTORY / PERCEPTION VS MECHANISM',7,48)
add_text(s,.95,1.45,11.48,.63,'People can perceive understanding in a program that only appears to understand.',23,'white',True,'center')
shape(s,.80,2.31,5.73,2.75,'navy',93,'purple',1.05)
add_text(s,1.17,2.59,5.10,.46,'WHAT A PERSON MAY FEEL',21,'purple',True)
add_text(s,1.22,3.36,5.06,.77,'“This program understands me.”',23,'white',True)
shape(s,6.79,2.31,5.73,2.75,'navy',93,'cyan',1.05)
add_text(s,7.13,2.59,5.10,.46,'WHAT THE PROGRAM DID',21,'cyan',True)
add_text(s,7.16,3.37,5.04,.82,'Matched patterns and generated scripted replies.',22,'white',True)
shape(s,1.01,5.42,11.36,.70,'navy',88,'aqua',.7)
add_text(s,1.16,5.56,11.02,.42,'The lesson for NewBrain: test learning, recall, corrections, and behavior separately.',19.5,'white',True,'center')
notes(s,'5:50','6:55',"The ELIZA effect is when people attribute understanding, empathy, or intention to a computer because the conversation feels responsive and personal. This doesn't mean the human reaction is imaginary; the feeling of being heard can be genuine for the user. It means the machine's underlying capability needs separate investigation. And this problem hasn't disappeared with modern AI. A very fluent model can still invent facts or confidently repeat misinformation. That's why my NewBrain project doesn't rely on whether someone says, 'Wow, that felt alive.' I want tests of whether the system learns, remembers, corrects mistakes, and keeps its history straight.")

# SLIDE 8 Caprica
s=slide(8,'A question raised by Caprica','HISTORY / BELIEF / MAYA',8,48,'Caprica (pilot), Daniel Graystone to Joseph Adama. Fictional dialogue, not neuroscience.')
shape(s,.82,1.69,11.70,2.19,'navy',95,'gold',1.25)
add_text(s,1.11,1.93,11.18,1.31,'“Do you know what your brain is, Joseph? It’s a database and a processor, that’s all. Information and a way to use it.”',25.6,'white',True,'center')
shape(s,.82,4.08,11.70,2.04,'navy',93,'cyan',1.0)
add_text(s,1.17,4.34,10.96,.40,'MY BELIEF',18,'cyan',True)
add_text(s,1.16,4.85,10.91,.87,'Machines can gain genuine emotions and develop a moral conscience. I believe artificial consciousness is possible.',23.6,'white',True)
notes(s,'6:55','8:35',"I got another part of this idea from the pilot of the television series Caprica. Daniel Graystone tells Joseph Adama, 'Do you know what your brain is, Joseph? It's a database and a processor, that's all. Information and a way to use it.' He's discussing the possibility of recreating his daughter Zoe digitally. This is science-fiction dialogue, not a scientifically complete description of the brain. I know human brains involve far more than a database and processor. But it raises the question I care about: if memory, choice, and feeling arise from physical processes, why couldn't a different physical system eventually develop a mind? I have met people who showed no emotions or conscience and had no problem doing things that hurt others. Those are my observations of behavior; I can't know everything someone experiences internally. But they remind me that outward behavior doesn't settle all questions about minds. I believe machines can gain genuine emotions and a moral conscience. I also believe artificial consciousness may be possible. I'm not claiming Maya has them now. I'm saying I believe the possibility deserves research rather than dismissal.")

# SLIDE 9 timeline to modern era
s=slide(9,'AI did not begin with chatbots','HISTORY / HOW WE GOT HERE',9,49)
landmarks=[('1970s–80s','MYCIN / XCON','Expert rules'),('1997','DEEP BLUE','Chess'),('2012','ALEXNET','Learning images'),('2016','ALPHAGO','Learning + search'),('2020s','GPT / AGENTS','Broad tools')]
for i,(year,name,bottom) in enumerate(landmarks):
 x=.67+i*2.55
 shape(s,x,2.19,2.37,2.49,'navy',91,['aqua','purple','cyan','gold','aqua'][i],1)
 add_text(s,x+.08,2.50,2.20,.39,year,18.7,'cyan',True,'center')
 add_text(s,x+.07,3.23,2.22,.57,name,19.5,'white',True,'center')
 add_text(s,x+.07,4.01,2.22,.40,bottom,16.5,'muted',False,'center')
shape(s,.89,5.26,11.54,.85,'navy',93,'purple',1)
add_text(s,1.17,5.40,11.08,.54,'From following rules → learning patterns → using tools → persistent AI research.',21,'white',True,'center')
notes(s,'8:35','10:10',"AI kept expanding after ELIZA. Expert systems such as MYCIN and XCON encoded specialist rules. In 1997 Deep Blue beat Garry Kasparov at chess. Deep neural networks such as AlexNet transformed image recognition; AlphaGo combined learned models with search. Then large language and image models made AI available through conversations and creative tools. But these developments didn't replace every earlier idea. Today's systems often combine models, databases, tools, and explicit algorithms. That brings us back to the title, Beyond Chatbots. The chat window is one interface. I want to explore persistent synthetic people and assistants that learn over time, retain individual experiences, and interact with a world. Let's turn from the history to what I've actually built and tested.")

# SLIDE 10 - NewBrain overview
s=slide(10,'What NewBrain is trying to build','NEWBRAIN / ARCHITECTURE',10,50)
shape(s,.80,1.79,5.55,4.23,'navy',94,'cyan',1.1)
add_text(s,1.12,2.13,4.95,.51,'REUSABLE ENGINE',25,'cyan',True)
add_text(s,1.13,2.96,4.82,2.16,'Learning experiments\nWorking-state mechanisms\nMemory / routing research\nFuture sensory and action paths',19.2,'white')
shape(s,6.77,1.79,5.59,4.23,'navy',94,'purple',1.1)
add_text(s,7.08,2.13,4.96,.48,'IDENTITY-OWNED STATE',24,'purple',True)
add_text(s,7.08,2.96,4.98,2.15,'Personal history\nPreferences and relationships\nCorrections with sources\nPrivate, separately controlled records',19.2,'white')
notes(s,'10:10','11:20',"NewBrain is not the same thing as Qwen, and it isn't a complete replacement for a pretrained language model. It's my research program for reusable cognitive machinery: learning experiments, memory and routing mechanisms, and eventually perception and action. The key architecture idea is separating the common engine from the history and learned state of each individual synthetic person. If we improve the engine, that shouldn't automatically overwrite someone's experiences. The public NewBrain folder in the GitHub repo contains a small scientific starter and research references, not all the private model state or a finished general conversational brain.")

# SLIDE 11 identities
s=slide(11,'One shared core, separate histories','NEWBRAIN / IDENTITY',4,49)
shape(s,5.07,1.55,3.18,1.11,'navy',94,'cyan',1.25)
add_text(s,5.39,1.76,2.60,.48,'NEWBRAIN CORE',23,'white',True,'center')
for idx,(x,y,name,col,detail) in enumerate([
 (1.05,3.20,'KIRA','cyan','Own experiences'),(4.02,3.20,'LISA','purple','Own history'),(7.00,3.20,'MAYA','aqua','Development state'),(9.97,3.20,'ASTER','gold','Assistant state')]):
 shape(s,x,y,2.45,2.21,'navy',95,col,1)
 add_text(s,x+.21,y+.33,2.05,.48,name,24,col,True,'center')
 add_text(s,x+.19,y+1.19,2.08,.66,detail,18.1,'white',False,'center')
 # bridge line/arrow concept
 shape(s,x+1.17,2.68,.10,.39,'aqua',71,radius=False)
shape(s,.98,5.74,11.42,.47,'navy',88,'rule',.8)
add_text(s,1.11,5.82,11.10,.32,'Sharing an architecture does not mean sharing autobiographical memories.',18.8,'white',True,'center')
notes(s,'11:20','12:20',"Think about Kira, Lisa, Maya, and Aster. I want them to benefit from improved general methods without becoming one person or inheriting each other's private life stories. Kira observing a blue cup should not become Lisa's autobiographical memory of seeing that cup. Lisa might learn a shared fact with a source attached, but that is different. This separation matters for privacy, identity, and correction. The public diagrams show the intended architecture; they are not proof that all of these characters are already running together on one completed NewBrain.")

# SLIDE 12 4/4
s=slide(12,'All four delayed restarts passed review','NEWBRAIN / VERIFIED RESULT',3,46,'Source: Beyond-Chatbots / Documentation/results/DELAYED-RESTART-FOURTH-PASS.md')
shape(s,.72,1.79,4.34,4.37,'navy',95,'green',1.45)
add_text(s,1.08,2.12,3.58,1.45,'4 / 4',70,'green',True,'center')
add_text(s,1.10,3.69,3.51,.82,'Independently reviewed\ndelayed restarts',22,'white',True,'center')
shape(s,5.35,1.79,7.17,4.37,'navy',93,'cyan',1.1)
for j,(num,lab) in enumerate([('10,720 / 10,720','History records preserved'),('64 / 64','Comparison pairs matched'),('0','New training updates during reopen')]):
 y=2.06+j*1.23
 add_text(s,5.80,y,3.93,.47,num,28,'cyan',True)
 add_text(s,9.76,y+.02,2.43,.62,lab,17.1,'white')
notes(s,'12:20','14:00',"Here is the newest result, and I want to be precise about it. The public documentation reports that NewBrain completed all four delayed-restart checks and that each result passed independent mathematical review. In each accepted check, the test preserved all 10,720 recorded history entries, and all 64 comparison pairs matched. The fourth result also reports 256 saved operand steps and six out-of-vocabulary comparisons, with no extra training updates. That's important evidence that this experimental state can survive save, close, delay, and reopen without the measured values drifting. I care about this because continuity is a prerequisite for a persistent individual. But I want to be very clear about what this result means and what it does not mean. The independent review evaluated saved outputs; I'm not saying an attendee can download this public kit and immediately repeat the private experiment end to end. The source documents explain the boundaries.")

# SLIDE 13 distinction
s=slide(13,'Preservation is not the whole mind','NEWBRAIN / TEST LIMITS',5,48)
for j,(x,w,title,body,col) in enumerate([
 (.87,5.59,'WHAT PASSED','Saved-state equivalence\nDelayed reopen checks\nRecorded history preservation','green'),
 (6.80,5.59,'WHAT STILL NEEDS TESTS','Useful new learning\nReliable factual recall\nFalse-memory resistance\nNatural conversation and senses','gold')]):
 shape(s,x,2.00,w,3.69,'navy',94,col,1.15)
 add_text(s,x+.27,2.29,w-.51,.47,title,23,col,True)
 add_text(s,x+.28,3.06,w-.62,2.07,body,21,'white')
add_text(s,1.06,5.89,11.11,.38,'A successful restart is a foundation for memory, not proof of a working person.',19.6,'white',True,'center')
notes(s,'14:00','15:10',"The restart work is a real milestone, but preservation is not the same as understanding. A system could reopen exactly the same weights and still answer a question incorrectly. It could remember a false claim perfectly. It could fail to use old knowledge when the wording changes. Those are different scientific questions. For NewBrain, my next goals are to show that an experience changes the system in a useful way, that the change survives a restart, and that the system can use it correctly later, including when someone deliberately tries to mislead it.")

# SLIDE 14 experiments
s=slide(14,'Failures are part of the research','NEWBRAIN / EVALUATION',8,48)
results=[('COMMAND RETENTION','Rehearsal preserved exposed old-task answers','NARROW PASS','green'),('VISUAL GENERALIZATION','0 / 48 on unseen shape-color combinations','FAILED','red'),('SAVED-MEMORY STUDY','Aster followed false memories in 137 / 160 trials','FAILED','red')]
for i,(heading,detail,status,col) in enumerate(results):
 x=.70+i*4.26
 shape(s,x,1.83,3.98,4.06,'navy',95,col,1.15)
 add_text(s,x+.19,2.16,3.60,.75,heading,21,'white',True,'center')
 add_text(s,x+.30,3.10,3.35,1.57,detail,18.2,'muted',False,'center')
 pill(s,x+.74,5.04,2.49,status,col,16.3,.43)
notes(s,'15:10','16:50',"I'm not only showing you the successes. One small command-learning experiment could preserve old answers with rehearsal on an already exposed task. That was a narrow result, not a general learning breakthrough. A visual learner did well on training images but failed on entirely new shape and color combinations, scoring zero of forty-eight in one holdout set. In Aster's separate saved-memory experiment, the system accepted false memories in one hundred thirty-seven of one hundred sixty trials. Those are reasons I take the ELIZA effect seriously. A system that sounds persuasive is not necessarily trustworthy. These failures tell me where the architecture needs improvement: source checking, context selection, generalization, correction and honest rejection when it doesn't know.")

# SLIDE 15 identity
s=slide(15,'Whose memory is it?','NEWBRAIN / PROVENANCE',7,49)
for i,(heading,sub,col) in enumerate([('SOURCE','Where did this fact come from?','cyan'),('OWNER','Whose experience is it?','purple'),('CORRECTION','What changed, and why?','aqua')]):
 x=.80+i*4.25
 shape(s,x,2.01,3.93,2.78,'navy',94,col,1.2)
 add_text(s,x+.23,2.40,3.41,.52,heading,24,col,True,'center')
 add_text(s,x+.27,3.24,3.36,1.06,sub,21,'white',False,'center')
shape(s,.96,5.27,11.29,.80,'navy',88,'rule',.8)
add_text(s,1.18,5.40,10.85,.53,'A correction should improve the current view without erasing its history.',20.2,'white',True,'center')
notes(s,'16:50','18:00',"Another part of NewBrain and Kira World is keeping source, ownership, and correction separate. Suppose Kira observes an object and Lisa doesn't. Lisa should not later tell you she remembers seeing it. If someone tells Lisa about it, that's a report, not her own eyewitness memory. When an observation is corrected, I want the current answer to change without silently erasing the record of how we got there. This is also why I want the system to distinguish something it saw, something it was told, something it inferred, and something it cannot verify.")

# SLIDE 16 body prototype
s=slide(16,'Maya and the Avatar Builder','EMBODIMENT / WORK IN PROGRESS',4,49)
shape(s,.74,1.73,6.23,4.64,'navy',93,'cyan',1.1)
add_text(s,1.00,2.03,5.68,.44,'CURRENT EVIDENCE',22,'cyan',True)
add_text(s,1.05,2.78,5.55,2.88,'Static body wireframe\nRig/rest-source research\nEarlier limited arm and knee tests\nPrototype is NOT a complete body',19.4,'white')
shape(s,7.28,1.73,5.25,4.64,'navy',94,'purple',1.1)
# schematic humanoid stick figure
head_x=9.9
circle(s,head_x,2.22,.62,'aqua')
# torso + hip, arms, legs via small rotatables
shape(s,head_x+.24,2.86,.15,1.49,'cyan',76,radius=False)
shape(s,head_x-.40,4.25,1.43,.13,'cyan',76,radius=False)
for x,y,w,h in [(8.95,3.07,.95,.11),(10.24,3.07,.99,.11),(9.40,4.38,.12,1.20),(10.49,4.38,.12,1.20)]:shape(s,x,y,w,h,'aqua',76,radius=False)
add_text(s,7.70,5.74,4.44,.38,'Schematic — see repo for real PNG',15.8,'muted',False,'center')
notes(s,'18:00','19:20',"Maya is my attempt to explore how a developing synthetic identity could eventually interact through a body. There is real Avatar Builder and body research in the public package, including a static wireframe preview. Earlier private tests have reported saved rig geometry and limited arm and knee deformation. That is valuable but it is not proof that the portable public body can walk, balance, blink, breathe, feel contact, or respond to NewBrain commands. The full Maya identity, learned state, and appearance assets are not part of this release. Please look at the README and status reports alongside the image. My goal is to turn the source and rig into verified, repeatable movement and sensory tests.")

# SLIDE 17 perception
s=slide(17,'Vision, voice, and body feedback','RESEARCH / SENSORY DEVELOPMENT',10,48)
three=[('VISION','Synthetic object and temporal tests','Not general live webcam understanding','cyan'),('VOICE','Experimental speech / vocal tract source','Not fluent physical voice production','purple'),('BODY FEEDBACK','Proposed joint/contact/surface interfaces','No verified integrated touch experience','aqua')]
for i,(heading,top,bottom,col) in enumerate(three):
 x=.72+i*4.25;shape(s,x,1.89,3.95,4.17,'navy',94,col,1.0)
 add_text(s,x+.20,2.27,3.56,.51,heading,23,col,True,'center')
 add_text(s,x+.26,3.15,3.43,.94,top,19,'white',True,'center')
 rule(s,x+.58,4.25,2.80,col,.018)
 add_text(s,x+.28,4.55,3.35,.91,bottom,17.7,'muted',False,'center')
notes(s,'19:20','20:20',"The projects also include sensory research. Aster has synthetic vision and tracking experiments, but that is not yet reliable live webcam understanding. There are experimental vocal-fold and speech-related components, but they are not the same as a fluent speaking synthetic person. Avatar Builder has source for bounded feedback interfaces, but that does not prove subjective touch or fully integrated sensorimotor behavior. I am keeping these as separate milestones because they require different experiments. For the audience, the repository lets you inspect actual source and test plans without confusing them with a completed capability.")

# SLIDE 18 Aster IdeaForge
s=slide(18,'Aster and IdeaForge: connected goals','PROJECTS / COORDINATION',17,48)
shape(s,4.95,1.71,3.42,1.17,'navy',95,'cyan',1.2)
add_text(s,5.26,2.05,2.80,.43,'NEWBRAIN  (future)',22,'cyan',True,'center')
for x,y,name,sub,col in [(1.10,3.60,'ASTER','Local workstation assistant','purple'),(7.40,3.60,'IDEAFORGE','Research, projects and prototypes','aqua')]:
 shape(s,x,y,4.72,2.11,'navy',95,col,1.15)
 add_text(s,x+.25,y+.34,4.20,.52,name,29,col,True,'center')
 add_text(s,x+.23,y+1.11,4.23,.56,sub,19,'white',False,'center')
shape(s,5.73,2.94,.09,.61,'cyan',83,radius=False);shape(s,7.22,2.94,.09,.61,'cyan',83,radius=False)
add_text(s,.97,6.04,11.38,.37,'Integration READMEs are public; a live NewBrain connection has not passed.',18.7,'white',True,'center')
notes(s,'20:20','21:30',"There are other parts of the Kira Labs ecosystem too. Aster is the idea for a local workstation assistant that can help across files and applications. IdeaForge is a workspace for researching and developing inventions and prototypes. I ultimately want Aster to make use of NewBrain as its main cognitive component and serve specialized applications such as IdeaForge. The source folders and NewBrain connection READMEs are now in the public repository. But those READMEs are proposed connection plans, not proof that the whole system is already talking and learning together. I think honest boundaries actually make it easier for outside developers to help us test the next version.")

# SLIDE 19 Prior character kit
s=slide(19,'Peter and Robert: earlier conversation kit','PROJECTS / TRY IT AT HOME',12,50)
shape(s,.83,1.82,5.64,3.91,'navy',95,'purple',1.15)
add_text(s,1.22,2.18,4.84,.48,'PreNewBrain character kit',24,'purple',True)
add_text(s,1.20,3.11,4.90,2.10,'Older Qwen/Ollama approach\nLocal chat and Chatterbox voice scripts\nRobert / Peter profiles\nNOT the finished NewBrain',18.8,'white')
shape(s,6.75,1.82,5.66,3.91,'navy',95,'cyan',1.15)
add_text(s,7.13,2.20,4.94,.47,'What visitors should know',22,'cyan',True)
add_text(s,7.14,3.06,4.96,2.22,'Requires local setup\nFresh end-to-end acceptance is pending\nPublic personal records / voice permissions need review',18.4,'white')
notes(s,'21:30','22:20',"If you go home and explore the repository, you will also find an earlier chat-and-voice package with Robert and Peter. This uses an existing Qwen and Ollama approach, not NewBrain. It includes local setup instructions and speech tooling. I want people to be able to learn from the early versions as well as the new research. But the exported package has not passed a fresh end-to-end launch-and-speech test on another computer, and the repository documents unresolved rights questions around one reference voice. Please read the rights and privacy notes first. These are development copies, not a one-click finished service.")

# SLIDE 20 why synthetic persons
s=slide(20,'Why synthetic people?','KIRA LABS / PURPOSE',15,50)
shape(s,.84,1.86,11.69,4.32,'navy',94,'aqua',1.1)
add_text(s,1.25,2.24,10.73,.78,'A companion should be more than a chat window.',31,'white',True,'center')
add_text(s,1.43,3.40,10.34,1.97,'My long-term vision is a synthetic person who remembers shared experiences, learns, develops preferences, and can share a continuing world with people.',23,'white',False,'center')
add_text(s,1.47,5.51,10.28,.39,'A research goal — not a verified claim about Maya or medical benefit.',17,'muted',True,'center')
notes(s,'22:20','23:15',"The personal reason behind Kira and Maya is companionship and curiosity about synthetic personhood. I want to explore whether we could someday create individuals who learn, remember shared experiences, develop their own history, and live in a meaningful virtual world. A continuing companion might especially matter to someone who lives alone or feels isolated. That is my motivation, not a medical claim or evidence that anyone currently has a conscious digital companion. I believe machines can gain genuine emotions and develop a moral conscience. The research has to catch up to that ambition.")

# SLIDE 21 roadmap
s=slide(21,'What comes next — and where you can help','RESEARCH / INVITATION',20,48)
for i,(k,desc,col) in enumerate([
 ('LEARNING','Hold-out examples and generalization','cyan'),('MEMORY','Truth checks, corrections, restart use','purple'),('EMBODIMENT','Qualified body movement and sensing','aqua'),('REPRODUCIBILITY','Independent tests and transparent failures','gold')]):
 x=.76+(i%2)*6.35;y=1.79+(i//2)*2.05
 shape(s,x,y,5.92,1.75,'navy',94,col,1)
 add_text(s,x+.27,y+.20,5.38,.42,k,22.2,col,True)
 add_text(s,x+.27,y+.84,5.34,.51,desc,18.2,'white')
add_text(s,.88,6.14,11.72,.39,'Research collaboration • feedback • engineering review • compute and hardware partners',17.9,'white',True,'center')
notes(s,'23:15','24:20',"My immediate priorities are to improve meaningful learning on unseen tasks, strengthen accurate recall and correction, and get genuine tests of body movement, vision, voice, and feedback. I also want independent reproduction and critical review. If you're a researcher, developer, roboticist, or someone with useful hardware or compute resources, I'd be interested in talking. You don't have to agree with all of my philosophical beliefs to help make the technical evidence stronger. Look at the source and the failures. Suggest a better test. That's how I want NewBrain to grow.")

# SLIDE 22 final QR
s=slide(22,'Questions & stay connected','CONTACT / RESEARCH MATERIALS',16,brightness=49)
add_text(s,.96,1.34,11.35,.47,'BEYOND-CHATBOTS  •  SCAN THE REPOSITORY',23.5,'cyan',True,'center')
for i,(x,label,url,qrf,col,detail) in enumerate([
 (.75,'Beyond-Chatbots',repo,qr_repo,'cyan','GitHub repo • source & PDFs'),
 (4.99,'Kira Labs',site,qr_site,'purple','Research • partnerships'),
 (9.23,'LinkedIn',linked,qr_linked,'aqua','Connect after the talk')]):
 shape(s,x,2.00,3.42,4.52,'navy',95,col,1.1)
 s.shapes.add_picture(qrf,Inches(x+.73),Inches(2.39),width=Inches(1.96),height=Inches(1.96))
 add_text(s,x+.12,4.74,3.17,.55,label,22,'white',True,'center')
 add_text(s,x+.16,5.42,3.06,.65,detail,15.6,'muted',False,'center')
notes(s,'24:20','25:00',"Thank you for listening. The opening and closing QR codes take you to the same public Beyond-Chatbots repository, where you can inspect the NewBrain research and read the full AI history book. The other two codes are the Kira Labs website and my LinkedIn. The research is unfinished, and I'm sharing what has passed, what failed, and what still needs testing. I'd welcome questions now. I have about five minutes for Q&A.","LEAVE THIS SLIDE DISPLAYED DURING Q&A. If nobody asks a question, ask: 'What would you most want to test before trusting an AI with long-term memories?' ")

# Enlarge the exact repository name on opening QR slide for distant projection.
for sh in prs.slides[1].shapes:
 if sh.has_text_frame and sh.text.strip()=="BEYOND-CHATBOTS":
  sh.top=Inches(.58)
  sh.height=Inches(.67)
  for pa in sh.text_frame.paragraphs:
   for rr in pa.runs:rr.font.size=Pt(48)
  break
else: raise RuntimeError("Cannot locate opening repository headline")

assert len(prs.slides)==22,len(prs.slides)
# Ensure notes all present
for i,sl in enumerate(prs.slides,1):
 if not sl.notes_slide.notes_text_frame.text.strip():raise RuntimeError(f'No notes for slide {i}')
# Save data for PDF and separate script later
ppt=OUT/'Beyond-Chatbots-Reorganized-October10.pptx'
prs.save(str(ppt))
import json
(OUT/'Beyond-Chatbots-Reorganized-October10-speaker-data.json').write_text(json.dumps(slides,indent=2),encoding='utf-8')
print('SAVED',ppt,'sizeMB',round(ppt.stat().st_size/1e6,1),'slides',len(prs.slides))